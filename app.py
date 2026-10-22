"""Lecture Notes -> Study Kit: summary, 10 MCQs and flashcards (Streamlit + Gemini Flash)."""

import io
import json
import os
import re

import streamlit as st
from pptx import Presentation
from pypdf import PdfReader

DEFAULT_MODEL = "gemini-3.5-flash"
MAX_CHARS = 60_000  # keeps the prompt well inside the model's limits
MIN_CHARS = 200  # below this the file is probably scanned images
NUM_QUESTIONS = 10


# --------------------------------------------------------------------------- #
# File reading
# --------------------------------------------------------------------------- #
def _shape_text(shape):
    """Collect text from a PPTX shape (text frames, tables, grouped shapes)."""
    parts = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        parts.append(shape.text_frame.text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    if hasattr(shape, "shapes"):  # group shape
        for child in shape.shapes:
            parts.extend(_shape_text(child))
    return [p for p in parts if p and p.strip()]


def extract_text(data: bytes, filename: str) -> str:
    """Return plain text from a PDF or PPTX file."""
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        return "\n\n".join(p for p in pages if p)
    if name.endswith(".pptx"):
        prs = Presentation(io.BytesIO(data))
        slides = []
        for i, slide in enumerate(prs.slides, start=1):
            parts = []
            for shape in slide.shapes:
                parts.extend(_shape_text(shape))
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"Notes: {notes}")
            if parts:
                slides.append(f"[Slide {i}]\n" + "\n".join(parts))
        return "\n\n".join(slides)
    raise ValueError("Unsupported file type. Please upload a PDF or PPTX file.")


# --------------------------------------------------------------------------- #
# Prompting and parsing
# --------------------------------------------------------------------------- #
def build_prompt(text: str, num_cards: int) -> str:
    return f"""You are an expert teacher. Using ONLY the lecture notes below, create study material.

Return ONLY valid JSON (no markdown, no commentary) in exactly this shape:
{{
  "summary": "A clear summary of the lecture in 150-250 words, in simple language.",
  "key_points": ["5 to 8 short key takeaways"],
  "quiz": [
    {{
      "question": "Question text",
      "options": ["Option A", "Option B", "Option C", "Option D"],
      "answer_index": 0,
      "explanation": "One or two sentences on why this answer is correct."
    }}
  ],
  "flashcards": [
    {{"front": "Term or question", "back": "Short, clear answer or definition"}}
  ]
}}

Rules:
- Exactly {NUM_QUESTIONS} items in "quiz", each with exactly 4 options and one correct answer.
- "answer_index" is the 0-based position (0-3) of the correct option. Vary the position.
- Distractors must be plausible but clearly wrong.
- Exactly {num_cards} items in "flashcards".
- Do not invent facts that are not supported by the notes.

LECTURE NOTES:
\"\"\"
{text}
\"\"\"
"""


def parse_json(raw: str) -> dict:
    """Parse model output as JSON, tolerating code fences or extra text."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        start, end = raw.find("{"), raw.rfind("}")
        if start != -1 and end > start:
            return json.loads(raw[start : end + 1])
        raise


def _answer_to_index(value, options):
    """Accept 0-3, 'B', or the answer text itself; return an int or None."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if 0 <= value < len(options) else None
    if isinstance(value, str):
        v = value.strip()
        if v.isdigit():
            return _answer_to_index(int(v), options)
        if len(v) == 1 and v.upper() in "ABCD":
            return "ABCD".index(v.upper())
        for i, opt in enumerate(options):
            if opt.strip().lower() == v.lower():
                return i
    return None


def validate_result(data: dict) -> dict:
    """Clean the model output and drop malformed items. Raises ValueError if unusable."""
    if not isinstance(data, dict):
        raise ValueError("Model did not return a JSON object.")

    quiz = []
    for item in data.get("quiz") or []:
        if not isinstance(item, dict):
            continue
        options = [str(o).strip() for o in (item.get("options") or []) if str(o).strip()]
        idx = _answer_to_index(item.get("answer_index", item.get("answer")), options)
        question = str(item.get("question", "")).strip()
        if question and len(options) == 4 and idx is not None:
            quiz.append(
                {
                    "question": question,
                    "options": options,
                    "answer_index": idx,
                    "explanation": str(item.get("explanation", "")).strip(),
                }
            )

    cards = []
    for item in data.get("flashcards") or []:
        if isinstance(item, dict):
            front, back = str(item.get("front", "")).strip(), str(item.get("back", "")).strip()
            if front and back:
                cards.append({"front": front, "back": back})

    summary = str(data.get("summary", "")).strip()
    points = [str(p).strip() for p in (data.get("key_points") or []) if str(p).strip()]

    if not summary and not quiz and not cards:
        raise ValueError("The model returned no usable content.")
    return {"summary": summary, "key_points": points, "quiz": quiz, "flashcards": cards}


def generate_study_kit(client, model: str, text: str, num_cards: int) -> dict:
    """Call Gemini and return validated study material. Retries once on bad JSON."""
    from google.genai import types  # imported here so the file reads/tests without the SDK

    config = types.GenerateContentConfig(
        response_mime_type="application/json", temperature=0.4
    )
    prompt = build_prompt(text[:MAX_CHARS], num_cards)
    last_error = None
    for _ in range(2):
        response = client.models.generate_content(model=model, contents=prompt, config=config)
        try:
            return validate_result(parse_json(response.text))
        except (ValueError, json.JSONDecodeError) as exc:
            last_error = exc
    raise ValueError(f"Could not read the AI response ({last_error}). Please try again.")


def to_markdown(result: dict, title: str) -> str:
    """Export the study kit as a Markdown document."""
    lines = [f"# Study Kit: {title}", "", "## Summary", result["summary"], ""]
    if result["key_points"]:
        lines += ["## Key Points"] + [f"- {p}" for p in result["key_points"]] + [""]
    lines.append("## Quiz")
    for i, q in enumerate(result["quiz"], start=1):
        lines.append(f"**{i}. {q['question']}**")
        for j, opt in enumerate(q["options"]):
            lines.append(f"- {'ABCD'[j]}. {opt}")
        lines.append(f"\nAnswer: {'ABCD'[q['answer_index']]}. {q['explanation']}\n")
    lines.append("## Flashcards")
    for c in result["flashcards"]:
        lines.append(f"- **{c['front']}** -> {c['back']}")
    return "\n".join(lines)


# --------------------------------------------------------------------------- #
# Streamlit UI
# --------------------------------------------------------------------------- #
def get_api_key() -> str:
    key = os.environ.get("GEMINI_API_KEY", "")
    try:
        key = st.secrets.get("GEMINI_API_KEY", key)  # raises if no secrets file exists
    except Exception:
        pass
    return key


def render_quiz(quiz):
    if not quiz:
        st.info("No quiz questions were generated. Try again.")
        return
    for i, q in enumerate(quiz):
        st.markdown(f"**Q{i + 1}. {q['question']}**")
        st.radio(
            "Choose one",
            range(len(q["options"])),
            format_func=lambda j, q=q: f"{'ABCD'[j]}. {q['options'][j]}",
            index=None,
            key=f"quiz_{i}",
            label_visibility="collapsed",
            disabled=st.session_state.get("submitted", False),
        )
        if st.session_state.get("submitted"):
            picked = st.session_state.get(f"quiz_{i}")
            right = q["answer_index"]
            if picked == right:
                st.success(f"Correct! {q['explanation']}")
            else:
                st.error(f"Correct answer: {'ABCD'[right]}. {q['options'][right]}. {q['explanation']}")
        st.divider()

    if not st.session_state.get("submitted"):
        if st.button("Submit quiz", type="primary"):
            st.session_state["submitted"] = True
            st.rerun()
    else:
        score = sum(st.session_state.get(f"quiz_{i}") == q["answer_index"] for i, q in enumerate(quiz))
        st.metric("Your score", f"{score} / {len(quiz)}")
        if st.button("Retake quiz"):
            for i in range(len(quiz)):
                st.session_state.pop(f"quiz_{i}", None)
            st.session_state["submitted"] = False
            st.rerun()


def main():
    st.set_page_config(page_title="Lecture Notes Study Kit", page_icon="📚", layout="centered")
    st.title("📚 Lecture Notes → Study Kit")
    st.caption("Upload a PDF or PPTX and get a summary, a 10-question quiz and flashcards.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_api_key() or st.text_input("Gemini API key", type="password")
        num_cards = st.slider("Number of flashcards", 5, 20, 10)
        model = st.text_input("Model", value=DEFAULT_MODEL)
        st.markdown("[Get a free API key](https://aistudio.google.com/app/apikey)")

    uploaded = st.file_uploader("Upload lecture notes", type=["pdf", "pptx"])

    if st.button("Generate study kit", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please add your Gemini API key in the sidebar or in Streamlit secrets.")
            st.stop()
        try:
            with st.spinner("Reading your file..."):
                text = extract_text(uploaded.getvalue(), uploaded.name)
            if len(text.strip()) < MIN_CHARS:
                st.error("Almost no text was found. The file may be scanned images; try a text-based PDF or PPTX.")
                st.stop()
            if len(text) > MAX_CHARS:
                st.warning(f"The file is long, so only the first {MAX_CHARS:,} characters were used.")
            from google import genai

            client = genai.Client(api_key=api_key)
            with st.spinner("Generating your study kit..."):
                result = generate_study_kit(client, model.strip() or DEFAULT_MODEL, text, num_cards)
            st.session_state.update(result=result, title=uploaded.name, submitted=False)
            for key in [k for k in st.session_state if k.startswith("quiz_")]:
                del st.session_state[key]
        except ValueError as exc:
            st.error(str(exc))
        except Exception as exc:  # network, quota, invalid key, corrupt file...
            st.error(f"Something went wrong: {exc}")

    result = st.session_state.get("result")
    if result:
        tab_summary, tab_quiz, tab_cards = st.tabs(["📝 Summary", "❓ Quiz", "🗂️ Flashcards"])
        with tab_summary:
            st.write(result["summary"])
            if result["key_points"]:
                st.subheader("Key points")
                for p in result["key_points"]:
                    st.markdown(f"- {p}")
        with tab_quiz:
            render_quiz(result["quiz"])
        with tab_cards:
            for c in result["flashcards"]:
                with st.expander(c["front"]):
                    st.write(c["back"])
        st.download_button(
            "Download study kit (.md)",
            to_markdown(result, st.session_state.get("title", "lecture")),
            file_name="study_kit.md",
            mime="text/markdown",
        )


if __name__ == "__main__":
    main()
